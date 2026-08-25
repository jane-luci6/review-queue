#!/usr/bin/env python3
"""Build the preliminary feature list for the new LUCI release as a Word doc.

Content source: Aug 25, 2026 product conversation (Richard, Nick).
Nothing internal-only (Hub, control-engine codenames, tunnel vendor) appears in the output.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path.home() / "Desktop" / "New LUCI - Preliminary Feature List.docx"

NAVY = RGBColor(0x35, 0x4F, 0x5C)
MINT_DARK = RGBColor(0x2B, 0x9E, 0x80)
INK = RGBColor(0x1F, 0x2C, 0x33)
MUTED = RGBColor(0x5C, 0x6B, 0x74)

HEAD_FONT = "Arial"
BODY_FONT = "Calibri"


def set_font(run, name=BODY_FONT, size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def para(doc, text="", size=11, bold=False, color=INK, italic=False,
         space_after=8, space_before=0, style=None, name=BODY_FONT):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        set_font(p.add_run(text), name=name, size=size, bold=bold, color=color, italic=italic)
    return p


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text.upper()), name=HEAD_FONT, size=13, bold=True, color=MINT_DARK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(10)
    set_font(rule.add_run("\u2014" * 28), name=HEAD_FONT, size=9, color=MINT_DARK)


def feature_heading(doc, text, number=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    label = f"{number}. {text}" if number else text
    set_font(p.add_run(label), name=HEAD_FONT, size=12, bold=True, color=NAVY)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        if isinstance(item, tuple):
            lead, rest = item
            set_font(p.add_run(lead), size=11, bold=True, color=INK)
            set_font(p.add_run(rest), size=11, color=INK)
        else:
            set_font(p.add_run(item), size=11, color=INK)


doc = Document()
for name in ("Normal", "List Bullet"):
    style = doc.styles[name]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    style.font.color.rgb = INK

for section in doc.sections:
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(58)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(2)
set_font(title.add_run("New LUCI \u2014 Preliminary Feature List"),
         name=HEAD_FONT, size=20, bold=True, color=NAVY)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(2)
set_font(sub.add_run("Features proposed for the launch campaign"),
         name=BODY_FONT, size=12, color=MUTED)

# ── Primary ──────────────────────────────────────────────────────────────────
section_heading(doc, "Primary Features")
para(doc, "The features we lead with \u2014 the ones we show repeatedly across the campaign.",
     size=11, color=MUTED, italic=True, space_after=6)

feature_heading(doc, "Property-branded interface", 1)
bullets(doc, [
    "Sign-in and splash screens carry the property\u2019s own photography and marks.",
    "Light or dark mode, with curated themes matched to the property\u2019s colors and typography chosen from an approved set.",
    "Background image treatment \u2014 blur, opacity, and positioning \u2014 is adjustable.",
    "Endpoint status colors arrive as operators already know them: green for healthy, yellow for attention, red for powered off. A property can change them across the whole install.",
    "LUCI can design the theme for a property, or their own team can build it and adjust it later.",
])

feature_heading(doc, "Venue panels", 2)
bullets(doc, [
    "Control lives in the venue. A tablet in the space puts control in the hands of the person who runs it \u2014 an events manager at a ballroom panel, a guest at a pool cabana panel.",
    "Assign a panel to a venue and it adopts every endpoint already assigned to that venue. No hand-built device list.",
    "Control is localized and scoped: the panel shows that one space, and only the controls the administrator allows \u2014 volume but not power, source but not volume, whatever fits the room.",
    "The person at the panel never sees the main LUCI application and cannot reach anything outside the designated space. A public-facing panel can be opened without a PIN; a staff panel can require one.",
    "Presets can be made available at the panel, so a room can be set to a saved look without leaving the space.",
    "A panel can also be scoped across several venues \u2014 a tablet that travels with the manager who covers four ballrooms.",
])

feature_heading(doc, "Staging", 3)
bullets(doc, [
    "Build the next look while the current one is still running: choose screens, set sources, set volumes, set content. Nothing changes in the room until Apply.",
    "Apply on cue. The change lands at the moment the event actually starts, not the moment it was scheduled to.",
    "Save any staged set as a preset to run it again. Give a preset a time and it becomes a scheduled preset.",
    "Scroll forward and back through days across new chart views to see what already ran and what is coming.",
    "Look at one display and see every preset and schedule due to touch it, and when.",
])

feature_heading(doc, "Audio group control", 4)
bullets(doc, [
    "A group slider moves several zones at once and keeps their balance \u2014 the dining room stays quieter than the bar, the patio stays louder.",
    "Move a group up or down by a step, set every zone in the group to the same level, or move a single zone on its own.",
    "A day spent tuning zones survives the first time someone reaches for a group control.",
])

feature_heading(doc, "Audit trails", 5)
bullets(doc, [
    "Every action is recorded against a person and a time: power, source, volume, mute, channel changes, schedule runs, overrides, and configuration changes.",
    "Each entry shows what triggered it \u2014 a person, a preset, or a schedule.",
    "Search by device to see everything that changed it, or by user to see everything one person did.",
    "Filter by user and action type, and download results as a CSV for reporting.",
    "Changes made outside LUCI \u2014 someone picking up a remote and switching a display off \u2014 are noticed and recorded for most third-party device types, and appear on the map.",
    "Device incidents open when a device stops responding and close when it comes back, so a fault has a duration instead of an anecdote.",
])

feature_heading(doc, "In-product support", 6)
bullets(doc, [
    "Raise a support request from inside LUCI, on the thing that is actually wrong \u2014 a device, an incident, an error.",
    "The request arrives with its context attached: what happened, what changed, and the relevant logs.",
    "Requests are opened by a person, not fired automatically by every device hiccup, so the property controls what escalates.",
])

# ── Secondary ────────────────────────────────────────────────────────────────
section_heading(doc, "Secondary Features")
para(doc, "Worth naming in supporting materials, but not what the campaign leads with.",
     size=11, color=MUTED, italic=True, space_after=6)

bullets(doc, [
    ("Live map accuracy. ", "The map reflects the real state of the floor through a live connection to devices rather than periodic polling, so what is on screen matches what is in the room."),
    ("Add any endpoint from LUCI. ", "Every device type can be added through the interface. Add one from the map or twenty at once with addresses incrementing, and get a warning when a record already exists."),
    ("Sign-in and session control. ", "Users sign in by email or PIN, or with the Microsoft credentials they already use at work through Entra ID. Administrators can sign a user out of the system and post a site-wide message banner to everyone using LUCI."),
    ("Video wall layout sync. ", "Wall layouts and window assignments are pulled from the video wall processor with per-pixel geometry instead of being rebuilt by hand. Multi-window templates can be assigned to a player and changed live."),
    ("Central display model catalog. ", "New hospitality display models are added centrally and pushed out to the property, so buying a new model no longer means editing configuration files on site."),
])

# ── IT ───────────────────────────────────────────────────────────────────────
section_heading(doc, "IT Technical Details")
para(doc, "The same features, with the detail an IT or A/V lead will ask for \u2014 not a separate product story.",
     size=11, color=MUTED, italic=True, space_after=6)

feature_heading(doc, "Venue panels")
bullets(doc, [
    "A panel is a scoped control surface tied to a point in the site hierarchy \u2014 site, building, floor, venue.",
    "On creation it inherits every endpoint assigned to that venue. Administrators then narrow it per control surface (power, source, volume, mute) rather than per device.",
    "The panel loads a pared-down interface with no route into the administrative application, and returns to its paired configuration on every boot.",
])

feature_heading(doc, "Staging, presets, and schedules")
bullets(doc, [
    "A staged set, a preset, and a scheduled preset are the same object at different stages: a set of commands, optionally named, optionally timed.",
    "The staged set is visible as a roster before anything is sent, and can be cleared or edited item by item. It can also be seeded from what the room is doing right now.",
    "Editing a saved preset shows the difference against the saved values before the change is committed.",
    "A scheduled preset can carry a lockout window, enforced at the control engine, so the devices it touched cannot be changed for a set period. A lock indicator appears on the device, and administrators can override.",
])

feature_heading(doc, "Audit trails and device incidents")
bullets(doc, [
    "Commands sent to devices are grouped by the event that triggered them, so one entry resolves to a user, a preset, or a schedule.",
    "Administrative actions such as permission changes are recorded alongside device commands.",
    "Visibility of out-of-band changes depends on what each device driver reports. Devices that do not report their own state changes will not surface them.",
    "A device incident opens automatically when a device stops responding and closes when it recovers, producing measured downtime.",
    "Incidents are distinct from support tickets. A person decides which incident becomes a ticket.",
])

feature_heading(doc, "Endpoint management")
bullets(doc, [
    "All device types can be created through the interface, including LUCI-supplied hardware, with bulk creation, incrementing addresses, and duplicate detection.",
])

feature_heading(doc, "Identity and sessions")
bullets(doc, [
    "Federated sign-in uses a standard Microsoft Entra ID app registration with a callback to the property\u2019s LUCI server. A profile is created on first sign-in, and a PIN can be assigned for floor use.",
    "Administrators can terminate active sessions and communicate to all signed-in users through a site-wide banner.",
])

feature_heading(doc, "Additional detail")
para(doc, "Items not on the promoted list, but the ones an IT audience will care about most.",
     size=11, color=MUTED, italic=True, space_after=6)
bullets(doc, [
    ("Private tunnel. ", "Today a LUCI system on property needs a long list of outbound destinations opened by the property\u2019s IT team. That consolidates into a single encrypted outbound connection between the on-property system and LUCI. Credentials no longer sit unrotated on the local machine \u2014 keys rotate on a schedule. For IT it is one paired connection to review instead of a list of exceptions, with no standing inbound access to the property network."),
    ("Scoped diagnostic capture. ", "Debug-level logging can be aimed at a single endpoint, driver, or module for a defined window \u2014 capture three minutes while reproducing a fault \u2014 instead of turning up logging across the system. The output can be reviewed on site or sent to LUCI with a support request."),
])

for p in doc.paragraphs:
    if p.alignment is None:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.save(OUT)
print(f"Wrote {OUT}")
