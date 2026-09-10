#!/usr/bin/env python3
"""Build the New LUCI net-new marketing campaign plan as a Word doc.

Separate from the existing-customer GTM brief. This one is for prospecting:
warm leads first (had a demo or requested one), then cold leads. Includes a review
cycle to update existing materials as-needed for the new release.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

OUT = Path.home() / "Desktop" / "New LUCI - Net New Campaign Plan.docx"

NAVY = RGBColor(0x35, 0x4F, 0x5C)
MINT_DARK = RGBColor(0x2B, 0x9E, 0x80)
INK = RGBColor(0x1F, 0x2C, 0x33)
MUTED = RGBColor(0x5C, 0x6B, 0x74)

HEAD_FONT = "Arial"
BODY_FONT = "Calibri"


def set_font(run, name=BODY_FONT, size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def para(doc, text="", size=11, bold=False, color=INK, italic=False,
         space_after=8, space_before=0, name=BODY_FONT):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        set_font(p.add_run(text), name=name, size=size, bold=bold, color=color, italic=italic)
    return p


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text.upper()), name=HEAD_FONT, size=13, bold=True, color=MINT_DARK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(10)
    set_font(rule.add_run("\u2014" * 28), name=HEAD_FONT, size=9, color=MINT_DARK)


def sub_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), name=HEAD_FONT, size=12, bold=True, color=NAVY)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        if isinstance(item, tuple):
            lead, rest = item
            set_font(p.add_run(lead), size=11, bold=True, color=INK)
            set_font(p.add_run(rest), size=11, color=INK)
        else:
            set_font(p.add_run(item), size=11, color=INK)


doc = Document()
for name in ("Normal", "List Bullet"):
    style = doc.styles[name]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    style.font.color.rgb = INK

for section in doc.sections:
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(58)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(2)
set_font(title.add_run("New LUCI \u2014 Net New Campaign Plan"),
        name=HEAD_FONT, size=20, bold=True, color=NAVY)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(2)
set_font(sub.add_run("Prospecting campaign \u2014 separate from the existing-customer GTM brief"),
        name=BODY_FONT, size=12, color=MUTED)

# ── 1. LUCI messaging ───────────────────────────────────────
section_heading(doc, "LUCI messaging")
para(doc, "This campaign uses LUCI\u2019s established messaging. It does not reinvent it.",
     size=11, color=MUTED, italic=True, space_after=6)
bullets(doc, [
    ("Brand promise: ", "The Orchestration Engine for Enterprise Multimedia."),
    ("Sub-tagline: ", "One interface to control, automate, and execute the entire guest experience."),
    ("Pillars: ", "Complete visibility and control; align teams; activate the guest experience; invest in A/V that scales and improves."),
])

# ── 2. Channels to target ─────────────────────────────────
section_heading(doc, "Channels to target")
bullets(doc, [
    ("LinkedIn. ", "Awareness, thought leadership, and feature highlights aimed at prospects."),
    ("Website. ", "SEO, feature pages, and a demo request call to action."),
    ("Outbound email. ", "Warm leads first, then cold."),
    ("Sales follow-up. ", "Direct outreach to warm leads who have seen LUCI before."),
    ("The Signal newsletter. ", "If a prospecting-appropriate issue runs. Open decision."),
    ("Industry events and partnerships. ", "Lower priority. Evaluate quarterly."),
])

# ── 3. Audience sequencing ─────────────────────────────────
section_heading(doc, "Audience sequencing")
para(doc, "Warm leads first, then cold.",
     size=11, color=MUTED, italic=True, space_after=6)
bullets(doc, [
    ("Phase 1 \u2014 Warm leads. ", "Properties that have had a LUCI demo or requested one. Direct follow-up with what is new since they saw it."),
    ("Phase 2 \u2014 Cold leads. ", "Awareness through LinkedIn and website, then outbound outreach."),
])

# ── 4. Prospect nurture and automated email cycle ───────────
section_heading(doc, "Prospect nurture and automated email cycle")
para(doc, "The flow from lead-in to sales handoff. Sales enablement is the primary close lever; the nurture cycle warms leads until sales can take the conversation.",
     size=11, color=MUTED, italic=True, space_after=6)

sub_heading(doc, "Lead sources")
bullets(doc, [
    ("Website demo request. ", "Demo request form submission."),
    ("LinkedIn. ", "Organic inbound plus posts."),
    ("Partner referral. ", "Potentially a big channel. Partner-sourced leads may get a warmer handoff."),
    ("The Signal newsletter. ", "If a prospecting-appropriate issue runs. Open decision."),
    ("Mark\u2019s relationships and ongoing conversations. ", "The warmest pool. Manual, not automated."),
])

sub_heading(doc, "The flow")
bullets(doc, [
    ("1. Lead comes in \u2192 immediate auto-response. ", "Demo request confirmation, welcome, or partner intro. One CTA: what to expect next."),
    ("2. Automated nurture sequence. ", "Three to five emails over two to four weeks. Educational: orchestration thesis, embedded team, what LUCI runs. Use cases matched to the lead\u2019s industry. Proof: case studies and scope stats. One CTA per email."),
    ("3. Manual follow-up trigger. ", "Lead shows intent (opens, clicks, replies), matches a named-account target, or Mark already knows them. The trigger is a person, not an email."),
    ("4. Sales handoff. ", "Personal email from Mark, not automated. References the lead\u2019s context. Books the demo."),
    ("5. Ongoing nurture. ", "Quarterly check-in for non-responders. New case study or newsletter drop. Industry-specific content. Re-engage when a trigger fires."),
])

sub_heading(doc, "Email assets to build")
bullets(doc, [
    ("Immediate auto-response. ", "One per lead source (demo request, content download, partner referral)."),
    ("Nurture emails. ", "Three to five, sequenced. Each tied to one CTA."),
    ("Sales handoff template. ", "For Mark\u2019s personal follow-up, with context placeholders."),
    ("Ongoing nurture emails. ", "Quarterly re-engagement for non-responders."),
])

# ── 5. Review cycle ─────────────────────────────────────────
section_heading(doc, "Review cycle: existing materials to update")
para(doc, "Review for the new release and notable feature changes. Do not rebuild from scratch.",
     size=11, color=MUTED, italic=True, space_after=6)
bullets(doc, [
    ("Website. ", "Feature pages, homepage, pillar pages."),
    ("Sales deck. ", "Update for the new release."),
    ("One-pager / spec sheet. ", "Update what LUCI is and is not."),
    ("Case studies. ", "If a pilot property can be named and cleared."),
    ("IT brief. ", "Update for technical buyers."),
    ("Other collateral. ", "Review as needed."),
])

# ── 6. Assets to build ───────────────────────────────────────
section_heading(doc, "Assets to build for the net new campaign")
para(doc, "Email assets are listed under the nurture cycle above.",
     size=11, color=MUTED, italic=True, space_after=6)
bullets(doc, [
    ("LinkedIn awareness posts. ", "Prospect-facing, not the same as the customer announcement."),
    ("Website demo request. ", "Landing page or CTA for booking a demo."),
    ("Sales talk track. ", "For prospecting conversations."),
])

# ── 7. Open decisions ─────────────────────────────────────
section_heading(doc, "Open decisions")
bullets(doc, [
    ("\u201cNewer, better LUCI\u201d as the campaign message. ", "Yes or no. Could anchor the prospecting campaign, but not decided."),
    ("The Signal newsletter. ", "Used for prospecting, or stays customer-only."),
    ("Paid channels. ", "If any budget, or organic only."),
    ("Pilot property for case study. ", "Whether a property can be named and cleared."),
    ("Demand-gen engine build. ", "Website, social, and partner channels are the feed into the nurture cycle. Scope and ownership for that build is a separate effort."),
])

doc.save(OUT)
print(f"Wrote {OUT}")
