#!/usr/bin/env python3
"""Fill the blanks in the MPSA master .docx from values (usually read from
Mike's uploaded proposal) and save a client-specific .docx.

This is the docx equivalent of the HTML templates' "populate in place" step:
the agent reads the proposal, extracts the values, and passes them here. The
master's legal boilerplate is never rewritten — only the underscore blanks are
replaced, run-by-run, preserving the document's formatting and styles.

Blanks filled (see scripts/build-mpsa-master.py for the header; this only
touches the body blanks):

  p#3  preamble:  Effective Date · Customer legal name · Customer entity type
                 · Customer principal-place-of-business address
  p#4  recital:  Facility address
  p#47 notices:  Customer notice address
  table 0 (signature block), per cell:
       Name · Title · Date   (cell 0 = LUCI signer, cell 1 = Customer signer)

Exhibit A (fee schedule) and Exhibit D (signed SOW) are attachments, not text
blanks — they are left as placeholders for Mike to attach in Word, or the
agent attaches the proposal fee schedule / signed SOW separately.

Usage:
  python3 scripts/fill-mpsa.py \\
    --out ~/Desktop/LUCI\\ Docs/<client>/clients/<client>-mpsa.docx \\
    --client "Elwha River Casino" \\
    --client-entity "a Washington tribal corporation" \\
    --client-address "123 Lower Elwha Pl, Port Angeles, WA 98363" \\
    --facility "Same as above" \\
    --notice-address "Attn: General Manager, 123 Lower Elwha Pl, Port Angeles, WA 98363" \\
    --effective-date "August 13, 2026" \\
    --luci-signer "Mike Epstein" --luci-title "CEO" \\
    --customer-signer "____" --customer-title "____"

Omit any value to leave its blank underscored (fill later in Word).
"""
import argparse, re
from pathlib import Path
from docx import Document

REPO = Path(__file__).resolve().parents[1]
DEFAULT_MASTER = REPO / 'ui_kits/sales/mpsa-vendor-protective.docx'
_UNDER = re.compile(r'_+')


def _replace_underscores(text, values):
    """Replace each _+ sequence in `text` with the next value from `values`,
    in order. Returns (new_text, count_consumed)."""
    if not values:
        return text, 0
    parts = _UNDER.split(text)
    seps = _UNDER.findall(text)
    if len(seps) > len(values):
        # more blanks than values left — fill what we have, keep rest as underscores
        pass
    out = [parts[0]]
    consumed = 0
    for i, sep in enumerate(seps):
        if consumed < len(values) and values[consumed] is not None and values[consumed] != '':
            out.append(values[consumed])
        else:
            out.append(sep)  # leave the underscores
        consumed += 1
        out.append(parts[i + 1])
    return ''.join(out), consumed


def _fill_paragraph(paragraph, values):
    """Replace underscore blanks across the paragraph's runs, in order.
    Returns the number of values consumed."""
    vi = 0
    remaining = list(values)
    for run in paragraph.runs:
        if '_' in run.text and remaining:
            new_text, used = _replace_underscores(run.text, remaining)
            run.text = new_text
            remaining = remaining[used:]
            vi += used
    return vi


def _fill_sig_cell(cell, name, title, date):
    """Fill the Name / Title / Date lines in a signature-block cell."""
    for p in cell.paragraphs:
        t = p.text.strip()
        if t.startswith('Name:'):
            _fill_paragraph(p, [name])
        elif t.startswith('Title:'):
            _fill_paragraph(p, [title])
        elif t.startswith('Date:'):
            _fill_paragraph(p, [date])


def fill(master, out, v):
    doc = Document(str(master))

    # p#3 preamble — 4 blanks in order.
    _fill_paragraph(doc.paragraphs[3], [
        v['effective_date'], v['client'], v['client_entity'], v['client_address'],
    ])
    # p#4 recital — facility address.
    _fill_paragraph(doc.paragraphs[4], [v['facility']])
    # p#47 notices — customer notice address.
    _fill_paragraph(doc.paragraphs[47], [v['notice_address']])

    # Signature block (table 0): cell 0 = LUCI, cell 1 = Customer.
    sig = doc.tables[0].rows[0].cells
    _fill_sig_cell(sig[0], v['luci_signer'], v['luci_title'], v['luci_date'])
    _fill_sig_cell(sig[1], v['customer_signer'], v['customer_title'], v['customer_date'])

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f'Filled MPSA → {out}')


def main():
    ap = argparse.ArgumentParser(description='Fill MPSA .docx blanks from values.')
    ap.add_argument('--master', default=str(DEFAULT_MASTER))
    ap.add_argument('--out', required=True)
    ap.add_argument('--client', default='')
    ap.add_argument('--client-entity', dest='client_entity', default='')
    ap.add_argument('--client-address', dest='client_address', default='')
    ap.add_argument('--facility', default='')
    ap.add_argument('--notice-address', dest='notice_address', default='')
    ap.add_argument('--effective-date', dest='effective_date', default='')
    ap.add_argument('--luci-signer', dest='luci_signer', default='')
    ap.add_argument('--luci-title', dest='luci_title', default='')
    ap.add_argument('--luci-date', dest='luci_date', default='')
    ap.add_argument('--customer-signer', dest='customer_signer', default='')
    ap.add_argument('--customer-title', dest='customer_title', default='')
    ap.add_argument('--customer-date', dest='customer_date', default='')
    a = ap.parse_args()
    v = {k: getattr(a, k) for k in (
        'client', 'client_entity', 'client_address', 'facility', 'notice_address',
        'effective_date', 'luci_signer', 'luci_title', 'luci_date',
        'customer_signer', 'customer_title', 'customer_date')}
    fill(a.master, a.out, v)


if __name__ == '__main__':
    main()
