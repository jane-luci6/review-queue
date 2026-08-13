#!/usr/bin/env python3
"""Add a professional header to the MPSA master .docx.

One-time master-setup step (re-run after re-copying the v6 .docx into
ui_kits/sales/). The footer is already professional (dynamic PAGE field);
this only adds the missing header:

    [LUCI logo]                              MASTER PURCHASE AND SERVICES AGREEMENT
    ──────────────────────────────────────────────────────────────────────────

The header is a borderless 2-col table (logo | right-aligned title) followed by
a thin bottom-border paragraph. Nothing else in the document is touched.

Usage:
  python3 scripts/build-mpsa-master.py            # default master path
  python3 scripts/build-mpsa-master.py <path.docx>
"""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = Path(__file__).resolve().parents[1]
DEFAULT_MASTER = REPO / 'ui_kits/sales/mpsa-vendor-protective.docx'
LOGO = REPO / 'assets/logos/luci-full-mintmark-blacktext.png'


def _set_cell_border_none(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        tcPr.append(el)


def _para_bottom_border(paragraph, sz=4, color='BFBFBF'):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_header(docx_path):
    doc = Document(str(docx_path))
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False

    # Clear any existing header paragraphs (start clean).
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)

    # 2-col table: logo (left) | title (right).
    table = header.add_table(rows=1, cols=2, width=section.page_width - section.left_margin - section.right_margin)
    table.autofit = True
    left, right = table.rows[0].cells
    _set_cell_border_none(left)
    _set_cell_border_none(right)

    # Left cell: LUCI logo, ~1.3" wide (keeps the mintmark legible at header scale).
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = lp.add_run()
    run.add_picture(str(LOGO), width=Inches(1.3))

    # Right cell: document title, small gray uppercase, right-aligned.
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    trun = rp.add_run('MASTER PURCHASE AND SERVICES AGREEMENT')
    trun.font.size = Pt(8)
    trun.font.bold = False
    trun.font.color.rgb = RGBColor(0x4A, 0x4A, 0x4A)
    # letterspacing via w:spacing val in twentieths of a point (~40 = 2pt)
    rPr = trun._element.get_or_add_rPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), '40')
    rPr.append(spacing)

    # Thin hairline under the header.
    rule = header.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(0)
    _para_bottom_border(rule, sz=4, color='BFBFBF')

    doc.save(str(docx_path))
    print(f'Added header to: {docx_path}')


if __name__ == '__main__':
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_MASTER
    if not path.is_file():
        sys.exit(f'Master not found: {path}')
    add_header(path)
