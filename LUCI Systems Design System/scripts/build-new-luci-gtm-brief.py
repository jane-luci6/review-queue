#!/usr/bin/env python3
"""Build the New LUCI go-to-market brief as a Word doc.

Short version: primary features, secondary features, and the assets to build for release date.
Pairs with the Preliminary Feature List.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

OUT = Path.home() / "Desktop" / "New LUCI - Go-to-Market Brief.docx"

NAVY = RGBColor(0x35, 0x4F, 0x5C)
MINT_DARK = RGBColor(0x2B, 0x9E, 0x80)
INK = RGBColor(0x1F, 0x2C, 0x33)
MUTED = RGBColor(0x5C, 0x6B, 0x74)

HEAD_FONT = "Arial"
BODY_FONT = "Calibri"


def set_font(run, name=BODY_FONT, size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def para(doc, text="", size=11, bold=False, color=INK, italic=False,
         space_after=8, space_before=0, name=BODY_FONT):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        set_font(p.add_run(text), name=name, size=size, bold=bold, color=color, italic=italic)
    return p


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text.upper()), name=HEAD_FONT, size=13, bold=True, color=MINT_DARK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(10)
    set_font(rule.add_run("\u2014" * 28), name=HEAD_FONT, size=9, color=MINT_DARK)


def feature_heading(doc, text, number=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    label = f"{number}. {text}" if number else text
    set_font(p.add_run(label), name=HEAD_FONT, size=12, bold=True, color=NAVY)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        if isinstance(item, tuple):
            lead, rest = item
            set_font(p.add_run(lead), size=11, bold=True, color=INK)
            set_font(p.add_run(rest), size=11, color=INK)
        else:
            set_font(p.add_run(item), size=11, color=INK)


doc = Document()
for name in ("Normal", "List Bullet"):
    style = doc.styles[name]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    style.font.color.rgb = INK

for section in doc.sections:
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(58)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(2)
set_font(title.add_run("New LUCI \u2014 Go-to-Market Brief"),
        name=HEAD_FONT, size=20, bold=True, color=NAVY)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(2)
set_font(sub.add_run("Pairs with the Preliminary Feature List"),
        name=BODY_FONT, size=12, color=MUTED)

# ── Primary features ───────────────────────────────────────────
section_heading(doc, "Primary Features")
para(doc, "The features we lead with.",
     size=11, color=MUTED, italic=True, space_after=6)

feature_heading(doc, "Venue panels", 1)
bullets(doc, [
    "Control lives in the venue. A tablet in the space puts control in the hands of the person who runs it \u2014 an events manager at a ballroom panel, a guest at a pool cabana panel.",
    "Assign a panel to a venue and it adopts every endpoint already assigned to that venue. No hand-built device list.",
    "Control is localized and scoped: the panel shows that one space, and only the controls the administrator allows.",
    "The person at the panel never sees the main LUCI application and cannot reach anything outside the designated space.",
    "Presets can be made available at the panel. A panel can also be scoped across several venues for a roaming administrator.",
])

feature_heading(doc, "Staging", 2)
bullets(doc, [
    "Build the next look while the current one is still running. Nothing changes in the room until Apply.",
    "Apply on cue, so the change lands at the moment the event actually starts.",
    "Save any staged set as a preset. Give a preset a time and it becomes a scheduled preset.",
    "Scroll forward and back through days across new chart views to see what already ran and what is coming.",
])

feature_heading(doc, "Audit trails and in-product support", 3)
bullets(doc, [
    "Every action is recorded against a person and a time, including what triggered it \u2014 a person, a preset, or a schedule.",
    "Search by device or user, filter by action type, and download the results as a CSV.",
    "Changes made outside LUCI are noticed and recorded for most third-party device types. Device incidents open when a device stops responding and close when it returns, so a fault has a duration.",
    "Raise a support request from the device, incident, or error itself. The request arrives with what happened, what changed, and the relevant logs already attached.",
    "A person chooses what to escalate, so the property stays in control.",
])

feature_heading(doc, "Customizable interface", 4)
bullets(doc, [
    "Sign-in and splash screens carry the property\u2019s own photography and marks.",
    "Light or dark mode, with curated themes matched to the property\u2019s colors and typography selected from an approved set.",
    "Background image blur, opacity, and positioning are adjustable.",
    "Endpoint status colors arrive as operators already know them and can be changed across the whole install.",
    "LUCI can design the theme for a property, or the property can build and adjust it.",
])

# ── Secondary features ──────────────────────────────────────────
section_heading(doc, "Secondary Features")
para(doc, "Worth naming in supporting materials, but not what the campaign leads with.",
     size=11, color=MUTED, italic=True, space_after=6)

bullets(doc, [
    ("Audio group control. ", "Move several zones proportionally so their tuned balance survives, set every zone to the same level, or move one zone alone."),
    ("Live map accuracy. ", "The map reflects the real state of the floor through a live connection to devices rather than periodic polling."),
    ("Add any endpoint from LUCI. ", "Every device type can be added through the interface. Add one from the map or twenty at once with addresses incrementing."),
    ("Sign-in and session control. ", "Email or PIN, or Microsoft Entra ID for federated sign-in. Administrators can sign a user out and post a site-wide banner."),
    ("Video wall layout sync. ", "Wall layouts and window assignments are pulled from the video wall processor with per-pixel geometry instead of rebuilt by hand."),
    ("Central display model catalog. ", "New hospitality display models are added centrally and pushed to the property, so buying a new model no longer means editing configuration files on site."),
])

# ── Assets to build ────────────────────────────────────────────
section_heading(doc, "Assets to build for release date")
para(doc, "The new version has already been teased in the customer newsletter. No other teasers are planned.",
     size=11, color=MUTED, italic=True, space_after=6)

bullets(doc, [
    ("Launch announcement. ", "LinkedIn post on release day."),
    ("Feature highlight posts. ", "One or two LinkedIn posts on primary features, using real screenshots."),
    ("Website update. ", "Feature page or homepage callout reflecting the new release."),
    ("One-pager / spec sheet. ", "Update what LUCI is and is not."),
    ("IT brief. ", "Deeper leave-behind for technical buyers, drawn from the IT Technical Details section of the Feature List."),
    ("Email to existing clients. ", "Upgrade logistics, what changes, what they need to do."),
    ("Internal enablement. ", "Pre-brief FDEs, support, and sales. Prepare a short FAQ and a Loom."),
    ("Newsletter follow-up. ", "Full article after launch. The teaser already ran."),
    ("Case study. ", "If a pilot property can be named and cleared."),
])

doc.save(OUT)
print(f"Wrote {OUT}")
